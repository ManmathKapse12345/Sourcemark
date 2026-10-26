"""HTTP surface and the shared contract with the Express API."""

from __future__ import annotations

import io
import json

import jsonschema
import pytest
from conftest import REPO
from fastapi.testclient import TestClient

from app.contract import AnalyzeRequest, AnalyzeResponse
from app.main import app

AUTH = {"X-Internal-Token": "test-internal-token-123"}
FIXTURES = REPO / "packages" / "contract" / "fixtures"
SCHEMAS = REPO / "packages" / "contract" / "schema"


@pytest.fixture
def client():
    return TestClient(app)


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def test_requires_internal_token(client):
    assert client.post("/analyze", json={}).status_code == 401
    assert client.post("/analyze", json={}, headers={"X-Internal-Token": "wrong"}).status_code == 401


def test_shared_fixtures_parse_with_python_models():
    AnalyzeRequest.model_validate(load(FIXTURES / "analyze-request.json"))
    AnalyzeResponse.model_validate(load(FIXTURES / "analyze-response.json"))


def test_real_response_matches_exported_json_schema(client):
    res = client.post("/analyze", json=load(FIXTURES / "analyze-request.json"), headers=AUTH)
    assert res.status_code == 200
    body = res.json()
    jsonschema.validate(body, load(SCHEMAS / "AnalyzeResponse.json"))
    assert {f["source"]["code"] for f in body["findings"]} == {"R-1", "R-4", "F-2"}


def test_request_schema_and_python_model_agree():
    schema = load(SCHEMAS / "AnalyzeRequest.json")
    fixture = load(FIXTURES / "analyze-request.json")
    jsonschema.validate(fixture, schema)
    fixture["options"]["surprise"] = True
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(fixture, schema)
    with pytest.raises(Exception):
        AnalyzeRequest.model_validate(fixture)


def test_parse_and_extract(client):
    md = (REPO / "eval" / "northfern-guidelines.md").read_bytes()
    res = client.post("/guidelines/extract", data={"prefix": "G"},
                      files={"file": ("guide.md", io.BytesIO(md), "text/markdown")}, headers=AUTH)
    assert res.status_code == 200 and len(res.json()["clauses"]) == 16
    jsonschema.validate(res.json(), load(SCHEMAS / "ExtractGuidelineResponse.json"))

    bad = client.post("/parse", files={"file": ("x.pdf", io.BytesIO(b"not a pdf"), "application/pdf")}, headers=AUTH)
    assert bad.status_code == 422


def test_docx_roundtrip(client):
    import docx

    doc = docx.Document()
    doc.add_heading("Claims", level=1)
    doc.add_paragraph("3.1 Do not call any product a miracle or a cure.")
    buf = io.BytesIO()
    doc.save(buf)
    res = client.post("/guidelines/extract", data={"prefix": "G"},
                      files={"file": ("g.docx", io.BytesIO(buf.getvalue()), "application/octet-stream")}, headers=AUTH)
    assert res.status_code == 200
    assert res.json()["clauses"] == [{"code": "G-3.1", "heading": "Claims", "text": "Do not call any product a miracle or a cure."}]


def test_embed_without_key_is_503(client):
    res = client.post("/embed", json={"texts": ["hello"], "task": "query"}, headers=AUTH)
    assert res.status_code == 503
